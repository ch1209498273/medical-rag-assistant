"""Build deterministic PDF/DOCX files for the fictional offline demo.

The generator intentionally has a small, explicit source allow-list.  The v2
patient-education roadmap remains a markdown planning artifact and is never
copied into ``demo/documents``.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import os
import re
import stat
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import pymupdf as fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE_DIR = PROJECT_ROOT / "demo" / "sources"
DEFAULT_DOCUMENT_DIR = PROJECT_ROOT / "demo" / "documents"
V1_SOURCE_STEMS = (
    "training-policy",
    "pre-dialysis-check",
    "learning-points",
    "dialysis-basics",
)
V2_ROADMAP_STEM = "patient-education-v2-roadmap"
DISCLAIMER = (
    "免责声明：本文件为虚构演示资料，不代表真实医疗制度或诊疗建议。"
    "仅用于展示本地文档检索、引用和安全拒答流程。"
)
FOOTER_TEXT = "虚构演示资料｜不代表真实医疗制度或诊疗建议"
FIXED_DATE = datetime(2026, 1, 1, 0, 0, 0, tzinfo=timezone.utc)
FIXED_ZIP_TIMESTAMP = (2020, 1, 1, 0, 0, 0)
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_PDF_ID_RE = re.compile(rb"/ID\[<([0-9A-Fa-f]{32})><([0-9A-Fa-f]{32})>\]")
_REPARSE_POINT_ATTRIBUTE = 0x400


def build_demo_documents(
    *,
    source_dir: Path = DEFAULT_SOURCE_DIR,
    document_dir: Path = DEFAULT_DOCUMENT_DIR,
) -> tuple[Path, ...]:
    """Render the four v1 sources as one PDF/DOCX pair each.

    Returns paths in stable filename order.  No source outside the explicit
    v1 allow-list is read, and the v2 roadmap is not emitted as a document.
    """

    source_root = _validate_source_root(source_dir)
    output_root = _validate_output_root(document_dir)
    output_root.mkdir(parents=True, exist_ok=True)
    generated: list[Path] = []
    for stem in V1_SOURCE_STEMS:
        source_path = source_root / f"{stem}.md"
        _validate_source_file(source_path, source_root)
        if not source_path.is_file():
            raise FileNotFoundError(f"missing fictional demo source: {stem}.md")
        source_text = source_path.read_text(encoding="utf-8")
        title, paragraphs = _normalise_source(source_text, stem)
        _render_docx(output_root / f"{stem}.docx", title, paragraphs)
        _render_pdf(output_root / f"{stem}.pdf", title, paragraphs)
        generated.extend(
            (output_root / f"{stem}.docx", output_root / f"{stem}.pdf")
        )

    # If a previous run was given a stale v2 output, remove only that exact
    # generated basename.  The normal output directory should never contain
    # it, and this keeps repeated local generation deterministic.
    for stale in (
        output_root / f"{V2_ROADMAP_STEM}.pdf",
        output_root / f"{V2_ROADMAP_STEM}.docx",
    ):
        if stale.exists():
            stale.unlink()
    return tuple(sorted(generated, key=lambda path: path.name))


def _validate_source_root(source_dir: Path) -> Path:
    candidate = _absolute_without_resolving(source_dir)
    if _contains_reparse_component(candidate):
        raise ValueError("approved demo source directory cannot contain a symlink or reparse point")
    try:
        source_root = candidate.resolve(strict=True)
        approved_root = DEFAULT_SOURCE_DIR.resolve(strict=True)
    except (OSError, RuntimeError) as error:
        raise ValueError("approved demo source directory is unavailable") from error
    try:
        source_root.relative_to(approved_root)
    except ValueError as error:
        raise ValueError("source directory must be within the approved demo source tree") from error
    if not source_root.is_dir():
        raise ValueError("approved demo source directory must be a directory")
    return source_root


def _validate_source_file(source_path: Path, source_root: Path) -> None:
    candidate = _absolute_without_resolving(source_path)
    if _contains_reparse_component(candidate):
        raise ValueError("approved demo source file cannot be a symlink or reparse point")
    try:
        resolved = candidate.resolve(strict=True)
    except (OSError, RuntimeError) as error:
        raise FileNotFoundError(f"missing fictional demo source: {source_path.name}") from error
    try:
        resolved.relative_to(source_root)
    except ValueError as error:
        raise ValueError("source file must remain within the approved demo source tree") from error


def _validate_output_root(document_dir: Path) -> Path:
    candidate = _absolute_without_resolving(document_dir)
    if _contains_reparse_component(candidate):
        raise ValueError("demo document output directory cannot contain a symlink or reparse point")
    try:
        return candidate.resolve(strict=False)
    except (OSError, RuntimeError) as error:
        raise ValueError("demo document output directory is unavailable") from error


def _absolute_without_resolving(path: Path) -> Path:
    value = Path(path)
    if not value.is_absolute():
        value = Path.cwd() / value
    return Path(os.path.abspath(os.fspath(value)))


def _contains_reparse_component(path: Path) -> bool:
    """Detect symlink/reparse components without following them first."""

    parts = path.parts
    current = Path(path.anchor) if path.anchor else Path()
    start = 1 if path.anchor else 0
    for part in parts[start:]:
        current /= part
        try:
            # ``lstat`` inspects the link/reparse point itself, so dangling
            # junctions are visible even though ``Path.exists()`` is false.
            details = current.lstat()
        except FileNotFoundError:
            # A normal not-yet-created component is safe to create later.
            continue
        except OSError as error:
            # Any other inspection failure is ambiguous; fail closed rather
            # than resolving or writing through an uninspectable component.
            raise ValueError("demo path cannot be inspected safely") from error
        if stat.S_ISLNK(details.st_mode):
            return True
        if getattr(details, "st_file_attributes", 0) & _REPARSE_POINT_ATTRIBUTE:
            return True
    return False


def _normalise_source(source_text: str, stem: str) -> tuple[str, list[str]]:
    title = stem
    paragraphs: list[str] = [DISCLAIMER]
    for raw_line in source_text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith(">"):
            # The generator owns the exact, stable disclaimer rendering.
            continue
        heading = _HEADING_RE.match(line)
        if heading:
            if len(heading.group(1)) == 1:
                title = heading.group(2).strip()
            else:
                paragraphs.append(heading.group(2).strip())
            continue
        if line.startswith("- "):
            line = line[2:].strip()
        paragraphs.append(line)
    if not any(text for text in paragraphs[1:]):
        raise ValueError(f"fictional demo source is empty: {stem}.md")
    return title, paragraphs


def _set_docx_properties(document: Document, title: str) -> None:
    properties = document.core_properties
    properties.author = "Public Demo"
    properties.category = "Fictional offline demonstration"
    properties.comments = "Fictional demo asset; not clinical guidance."
    properties.content_status = "Final"
    properties.created = FIXED_DATE
    properties.identifier = f"fictional-demo-{title}"
    properties.keywords = "fictional,demo,offline,hemodialysis"
    properties.language = "zh-CN"
    properties.last_modified_by = "Public Demo"
    properties.last_printed = FIXED_DATE
    properties.modified = FIXED_DATE
    properties.revision = 1
    properties.subject = DISCLAIMER
    properties.title = title
    properties.version = "1.0"


def _render_docx(path: Path, title: str, paragraphs: list[str]) -> None:
    document = Document()
    _set_docx_properties(document, title)
    section = document.sections[0]
    section.top_margin = Pt(48)
    section.bottom_margin = Pt(54)
    section.left_margin = Pt(54)
    section.right_margin = Pt(54)

    title_paragraph = document.add_paragraph(style="Title")
    title_paragraph.add_run(title)
    disclaimer = document.add_paragraph()
    disclaimer.add_run(DISCLAIMER).bold = True
    for index, text in enumerate(paragraphs[1:], start=1):
        paragraph = document.add_paragraph()
        if index == 1 or text.endswith(("流程", "检查", "要求", "边界", "说明", "登记与补修")):
            paragraph.style = "Heading 2"
        paragraph.add_run(text)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(FOOTER_TEXT)
    footer_run.font.size = Pt(8)

    buffer = io.BytesIO()
    document.save(buffer)
    _write_stable_docx(path, buffer.getvalue())


def _write_stable_docx(path: Path, raw_docx: bytes) -> None:
    """Normalise ZIP member order/timestamps so DOCX hashes are repeatable."""

    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    if temporary.exists():
        temporary.unlink()
    with zipfile.ZipFile(io.BytesIO(raw_docx), "r") as source, zipfile.ZipFile(
        temporary,
        "w",
        compression=zipfile.ZIP_DEFLATED,
        compresslevel=9,
    ) as target:
        for info in sorted(source.infolist(), key=lambda item: item.filename):
            stable = zipfile.ZipInfo(info.filename, FIXED_ZIP_TIMESTAMP)
            stable.compress_type = zipfile.ZIP_DEFLATED
            stable.create_system = 0
            stable.external_attr = 0
            stable.extra = b""
            stable.comment = b""
            target.writestr(stable, source.read(info.filename))
    os.replace(temporary, path)


def _render_pdf(path: Path, title: str, paragraphs: list[str]) -> None:
    document = fitz.open()
    document.set_metadata(
        {
            "format": "PDF 1.7",
            "title": title,
            "author": "Public Demo",
            "subject": DISCLAIMER,
            "keywords": "fictional,demo,offline,hemodialysis",
            "creator": "build_demo_documents.py",
            "producer": "PyMuPDF",
            "creationDate": "D:20260101000000Z",
            "modDate": "D:20260101000000Z",
        }
    )
    page = document.new_page(width=595, height=842)
    body = "\n\n".join((title, *paragraphs))
    body_rect = fitz.Rect(56, 52, 539, 770)
    spare_height = page.insert_textbox(
        body_rect,
        body,
        fontsize=11,
        lineheight=1.35,
        fontname="china-s",
        align=0,
    )
    if spare_height < 0:
        document.close()
        raise ValueError(f"fictional demo source does not fit on one page: {title}")
    page.insert_text(
        (56, 813),
        FOOTER_TEXT,
        fontsize=8,
        fontname="china-s",
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    if temporary.exists():
        temporary.unlink()
    document.save(temporary, garbage=4, clean=True, deflate=True)
    document.close()
    raw_pdf = temporary.read_bytes()
    stable_id = hashlib.sha256(f"fictional-demo-pdf:{title}".encode()).hexdigest()[:32].upper()
    stable_bytes, replacements = _PDF_ID_RE.subn(
        f"/ID[<{stable_id}><{stable_id}>]".encode("ascii"),
        raw_pdf,
    )
    if replacements != 1:
        temporary.unlink(missing_ok=True)
        raise ValueError(f"could not normalise generated PDF metadata: {title}")
    temporary.write_bytes(stable_bytes)
    os.replace(temporary, path)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source-dir",
        type=Path,
        default=DEFAULT_SOURCE_DIR,
        help="fictional markdown source directory (defaults to demo/sources)",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_DOCUMENT_DIR,
        help="generated document directory (defaults to demo/documents)",
    )
    args = parser.parse_args()
    paths = build_demo_documents(
        source_dir=args.source_dir,
        document_dir=args.output_dir,
    )
    for path in paths:
        print(f"{path.name} sha256={_sha256(path)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
