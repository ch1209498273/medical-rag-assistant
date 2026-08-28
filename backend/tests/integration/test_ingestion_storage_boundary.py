"""Black-box acceptance checks across parsing, chunking, and storage.

These tests use only generated fictional documents.  They intentionally stop at
the parser/chunker/storage boundary: the Task 5 ingestion service and cloud
embedding/OCR adapters are not part of this gate.
"""

from __future__ import annotations

from hashlib import sha256
from pathlib import Path

import fitz
from app.domain.models import ExtractedBlock, ExtractedDocument, SourceRef
from app.ingestion.chunker import chunk_document
from app.ingestion.extractors import extract_document
from app.storage.qdrant import QdrantLocalVectorStore
from app.storage.sqlite import SqliteDocumentRepository
from docx import Document


def _embedding(first_value: float = 1.0) -> list[float]:
    """Return a deterministic fictional 1024-dimensional embedding."""

    return [first_value] + [0.0] * 1023


def _file_sha256(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def _write_fictional_docx(path: Path, body: str) -> None:
    document = Document()
    document.add_heading("虚构透析护理制度", level=1)
    document.add_paragraph(body)
    document.save(path)


def _write_blank_pdf(path: Path) -> None:
    pdf = fitz.open()
    pdf.new_page()
    pdf.save(path)
    pdf.close()


def test_docx_chunks_publish_replace_and_survive_restart(tmp_path: Path) -> None:
    """A parsed policy chunk must remain locatable across version replacement."""

    document_path = tmp_path / "制度.docx"
    _write_fictional_docx(document_path, "旧版虚构条款：交接班前完成设备点检。")
    first_document = extract_document(document_path)

    assert first_document.needs_ocr is False
    assert first_document.blocks[0].source.heading_path == ("虚构透析护理制度",)
    first_chunks = chunk_document(first_document)
    assert first_chunks[0].source.paragraph_start == 2
    assert first_chunks[0].source.file_name == "制度.docx"

    database_path = tmp_path / "metadata.sqlite3"
    qdrant_path = tmp_path / "qdrant"
    repository = SqliteDocumentRepository(database_path)
    vector_store = QdrantLocalVectorStore(qdrant_path)
    try:
        first = repository.begin_index("制度.docx", _file_sha256(document_path))
        assert first.action == "index"
        vector_store.upsert(
            first.version_id,
            first_chunks,
            [_embedding() for _ in first_chunks],
        )
        repository.activate(first.version_id)

        first_hits = vector_store.search(
            _embedding(), repository.active_version_ids(), limit=5
        )
        assert [hit.chunk.text for hit in first_hits] == [
            "旧版虚构条款：交接班前完成设备点检。"
        ]
        assert first_hits[0].chunk.source.paragraph_start == 2

        unchanged = repository.begin_index("制度.docx", _file_sha256(document_path))
        assert unchanged.action == "skip"
        assert unchanged.version_id == first.version_id

        _write_fictional_docx(document_path, "新版虚构条款：交接班后完成记录复核。")
        second_document = extract_document(document_path)
        second_chunks = chunk_document(second_document)
        second = repository.begin_index("制度.docx", _file_sha256(document_path))
        assert second.action == "index"
        assert second.version_id != first.version_id

        vector_store.upsert(
            second.version_id,
            second_chunks,
            [_embedding() for _ in second_chunks],
        )
        repository.activate(second.version_id)
        assert repository.get_active("制度.docx").version_id == second.version_id
        assert repository.get_version(first.version_id).status == "superseded"

        active_hits = vector_store.search(
            _embedding(), repository.active_version_ids(), limit=5
        )
        assert [hit.chunk.text for hit in active_hits] == [
            "新版虚构条款：交接班后完成记录复核。"
        ]
    finally:
        vector_store.close()
        repository.close()

    restarted_repository = SqliteDocumentRepository(database_path)
    restarted_store = QdrantLocalVectorStore(qdrant_path)
    try:
        active_ids = restarted_repository.active_version_ids()
        assert active_ids == {second.version_id}
        restarted_hits = restarted_store.search(_embedding(), active_ids, limit=5)
        assert [hit.chunk.text for hit in restarted_hits] == [
            "新版虚构条款：交接班后完成记录复核。"
        ]
        assert restarted_store.search(_embedding(), set(), limit=5) == []
    finally:
        restarted_store.close()
        restarted_repository.close()


def test_scanned_pdf_ocr_lifecycle_recovers_and_preserves_active_on_failure(
    tmp_path: Path,
) -> None:
    """OCR recovery publishes only resumed content and failed OCR keeps old data."""

    scan_path = tmp_path / "扫描制度.pdf"
    _write_blank_pdf(scan_path)
    scanned_document = extract_document(scan_path)
    assert scanned_document.needs_ocr is True
    assert scanned_document.blocks == ()

    ocr_document = ExtractedDocument(
        file_name=scan_path.name,
        blocks=(
            ExtractedBlock(
                text="虚构 OCR 条款：扫描件完成识别后才允许发布。",
                source=SourceRef(file_name=scan_path.name, page=1, page_end=1),
            ),
        ),
    )
    ocr_chunks = chunk_document(ocr_document)

    database_path = tmp_path / "metadata.sqlite3"
    qdrant_path = tmp_path / "qdrant"
    repository = SqliteDocumentRepository(database_path)
    vector_store = QdrantLocalVectorStore(qdrant_path)
    try:
        old = repository.begin_index(scan_path.name, "fictional-old-hash")
        repository.activate(old.version_id)

        pending = repository.begin_index(
            scan_path.name, _file_sha256(scan_path), needs_ocr=True
        )
        assert pending.action == "needs_ocr"
        assert repository.get_active(scan_path.name).version_id == old.version_id

        resumed = repository.begin_ocr_indexing(pending.version_id)
        assert resumed.status == "indexing"
        vector_store.upsert(resumed.version_id, ocr_chunks, [_embedding()])
        repository.activate(resumed.version_id)
        assert repository.get_active(scan_path.name).version_id == resumed.version_id
        assert repository.get_version(old.version_id).status == "superseded"

        failed = repository.begin_index(
            scan_path.name, "fictional-failed-ocr-hash", needs_ocr=True
        )
        repository.fail(failed.version_id, "fictional OCR unavailable")
        assert repository.get_active(scan_path.name).version_id == resumed.version_id
        assert repository.get_version(failed.version_id).status == "failed"
    finally:
        vector_store.close()
        repository.close()
