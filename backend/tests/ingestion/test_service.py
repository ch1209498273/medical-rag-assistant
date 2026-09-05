"""TDD tests for safe incremental policy indexing.

All documents and model responses in this module are fictional and local.  The
tests intentionally inject a fake embedder/vector store so they never spend a
cloud token or read the user's enterprise source directory.
"""

from __future__ import annotations

import asyncio
import os
import threading
import time
from dataclasses import dataclass
from pathlib import Path

import fitz
import pytest
from app.errors import ProviderError
from app.storage.sqlite import SqliteDocumentRepository
from docx import Document


def _write_policy(path: Path, text: str = "虚构条款：交接班前完成设备点检。") -> None:
    document = Document()
    document.add_heading("虚构透析护理制度", level=1)
    document.add_paragraph(text)
    document.save(path)


def _write_blank_pdf(path: Path) -> None:
    pdf = fitz.open()
    pdf.new_page()
    pdf.save(path)
    pdf.close()


@dataclass
class FakeEmbedder:
    """Count calls and return deterministic vectors without network access."""

    calls: int = 0
    raise_on_call: Exception | None = None

    async def embed(self, texts: list[str]) -> list[list[float]]:
        self.calls += 1
        if self.raise_on_call is not None:
            raise self.raise_on_call
        return [[float(index + 1), 0.0] for index, _ in enumerate(texts)]


@dataclass
class SlowEmbedder(FakeEmbedder):
    delay: float = 0.05

    async def embed(self, texts: list[str]) -> list[list[float]]:
        self.calls += 1
        await asyncio.sleep(self.delay)
        return [[float(index + 1), 0.0] for index, _ in enumerate(texts)]


class FakeVectorStore:
    def __init__(self) -> None:
        self.upserts: list[tuple[str, list[object], list[list[float]]]] = []

    def upsert(self, version_id, chunks, vectors) -> None:
        self.upserts.append((version_id, list(chunks), [list(vector) for vector in vectors]))


class FailingVectorStore(FakeVectorStore):
    def __init__(self, *, cleanup_fails: bool = False) -> None:
        super().__init__()
        self.deleted: list[str] = []
        self.cleanup_fails = cleanup_fails

    def upsert(self, version_id, chunks, vectors) -> None:
        super().upsert(version_id, chunks, vectors)
        raise RuntimeError("fictional partial vector write failure")

    def delete_version(self, version_id: str) -> None:
        self.deleted.append(version_id)
        if self.cleanup_fails:
            raise RuntimeError("fictional cleanup failure")


class FakeOcr:
    async def ocr_page(self, image_bytes: bytes) -> str:
        assert image_bytes
        return "虚构 OCR 条款：扫描件完成识别后才允许发布。"


@dataclass
class BlockingSyncOcr:
    delay: float = 0.2
    calls: int = 0

    def ocr_page(self, image_bytes: bytes) -> str:
        import time

        self.calls += 1
        time.sleep(self.delay)
        return "虚构同步 OCR 条款"


@pytest.fixture
def repository(tmp_path: Path):
    instance = SqliteDocumentRepository(tmp_path / "metadata.sqlite3")
    try:
        yield instance
    finally:
        instance.close()


@pytest.fixture
def policy_path(tmp_path: Path) -> Path:
    path = tmp_path / "制度.docx"
    _write_policy(path)
    return path


@pytest.fixture
def service(tmp_path: Path, repository, embedder: FakeEmbedder, vector_store: FakeVectorStore):
    from app.ingestion.service import IngestionService

    return IngestionService(
        repository=repository,
        vector_store=vector_store,
        embedder=embedder,
        source_documents_dir=tmp_path,
    )


@pytest.fixture
def embedder() -> FakeEmbedder:
    return FakeEmbedder()


@pytest.fixture
def vector_store() -> FakeVectorStore:
    return FakeVectorStore()


@pytest.mark.asyncio
async def test_second_index_of_unchanged_file_skips_embedding(
    service, embedder: FakeEmbedder, policy_path: Path
):
    first = await service.index(policy_path)
    second = await service.index(policy_path)

    assert first.status == "active"
    assert second.status == "skipped"
    assert embedder.calls == 1


@pytest.mark.asyncio
async def test_embedding_failure_keeps_previous_version_active(
    service, repository, embedder: FakeEmbedder, policy_path: Path
):
    old = await service.index(policy_path)
    _write_policy(policy_path, "虚构新版条款：交接班后复核记录。")
    embedder.raise_on_call = ProviderError(
        "siliconflow", "embed", 503, True, "服务暂不可用"
    )

    failed = await service.index(policy_path)

    assert failed.status == "failed"
    assert repository.get_active(policy_path.name).version_id == old.version_id


@pytest.mark.asyncio
async def test_removed_file_is_deactivated_and_restored_same_hash_is_reembedded(
    service, repository, embedder: FakeEmbedder, policy_path: Path
):
    original = policy_path.read_bytes()
    await service.scan()

    policy_path.unlink()
    removed = await service.scan()

    assert removed.deactivated == [policy_path.name]
    assert repository.get_active(policy_path.name) is None

    policy_path.write_bytes(original)
    restored = await service.scan()

    assert restored.results[0].status == "active"
    assert embedder.calls == 2


@pytest.mark.asyncio
async def test_incomplete_directory_scan_does_not_deactivate_missing_documents(
    service, repository, policy_path: Path, monkeypatch
):
    await service.index(policy_path)

    original_iterdir = Path.iterdir

    def broken_iterdir(directory: Path):
        if directory == service.source_documents_dir:
            raise OSError("fictional directory enumeration failure")
        return original_iterdir(directory)

    monkeypatch.setattr(Path, "iterdir", broken_iterdir)

    with pytest.raises(OSError):
        await service.scan()

    assert repository.get_active(policy_path.name) is not None


@pytest.mark.asyncio
async def test_index_rejects_path_outside_configured_source_directory(
    service, tmp_path: Path
):
    outside = tmp_path.parent / "outside.docx"
    _write_policy(outside)

    with pytest.raises(ValueError, match="source directory"):
        await service.index(outside)


@pytest.mark.asyncio
async def test_pending_needs_ocr_can_resume_when_ocr_is_configured_later(
    tmp_path: Path, repository, embedder: FakeEmbedder, vector_store: FakeVectorStore
):
    from app.ingestion.service import IngestionService

    policy_path = tmp_path / "扫描制度.pdf"
    _write_blank_pdf(policy_path)
    without_ocr = IngestionService(
        repository, vector_store, embedder, tmp_path, ocr_client=None
    )

    first = await without_ocr.index(policy_path)
    assert first.status == "needs_ocr"

    with_ocr = IngestionService(
        repository, vector_store, embedder, tmp_path, ocr_client=FakeOcr()
    )
    resumed = await with_ocr.index(policy_path)

    assert resumed.status == "active"
    assert repository.get_active(policy_path.name).version_id == resumed.version_id
    assert embedder.calls == 1


@pytest.mark.asyncio
async def test_repeating_scan_without_ocr_keeps_pending_version(
    tmp_path: Path, repository, embedder: FakeEmbedder, vector_store: FakeVectorStore
):
    from app.ingestion.service import IngestionService

    policy_path = tmp_path / "扫描制度.pdf"
    _write_blank_pdf(policy_path)
    service = IngestionService(
        repository, vector_store, embedder, tmp_path, ocr_client=None
    )

    first = await service.index(policy_path)
    second = await service.index(policy_path)

    assert first.status == "needs_ocr"
    assert second.status == "needs_ocr"
    assert repository.get_version(first.version_id).status == "needs_ocr"
    assert embedder.calls == 0


@pytest.mark.asyncio
async def test_scan_failure_uses_safe_code_without_echoing_path_or_document_text(
    service, policy_path: Path, monkeypatch
):
    original_open = os.open

    def unreadable(path, flags, *args, **kwargs):
        if Path(path) == policy_path:
            raise OSError(f"permission denied: {path} 虚构正文泄漏")
        return original_open(path, flags, *args, **kwargs)

    monkeypatch.setattr("app.ingestion.service.os.open", unreadable)

    result = await service.scan()

    assert result.results[0].status == "failed"
    assert result.results[0].failure_reason == "SOURCE_FILE_UNREADABLE"
    assert str(policy_path) not in result.results[0].failure_reason
    assert "虚构正文" not in result.results[0].failure_reason


@pytest.mark.asyncio
async def test_oversized_file_fails_before_parsing_or_provider_call(
    tmp_path: Path, repository, vector_store: FakeVectorStore, embedder: FakeEmbedder
):
    from app.ingestion.service import IngestionService

    policy_path = tmp_path / "oversized.docx"
    policy_path.write_bytes(b"x" * 11)
    service = IngestionService(
        repository, vector_store, embedder, tmp_path, max_file_bytes=10
    )

    result = await service.index(policy_path)

    assert result.status == "failed"
    assert result.failure_reason == "DOCUMENT_TOO_LARGE"
    assert embedder.calls == 0
    assert repository.get_active(policy_path.name) is None


@pytest.mark.asyncio
async def test_pdf_page_limit_fails_before_ocr_or_embedding(
    tmp_path: Path, repository, vector_store: FakeVectorStore, embedder: FakeEmbedder
):
    from app.ingestion.service import IngestionService

    policy_path = tmp_path / "many-pages.pdf"
    pdf = fitz.open()
    pdf.new_page()
    pdf.new_page()
    pdf.save(policy_path)
    pdf.close()
    ocr = FakeOcr()
    service = IngestionService(
        repository,
        vector_store,
        embedder,
        tmp_path,
        ocr_client=ocr,
        max_pdf_pages=1,
    )

    result = await service.index(policy_path)

    assert result.failure_reason == "DOCUMENT_LIMIT_EXCEEDED"
    assert embedder.calls == 0
    assert repository.get_active(policy_path.name) is None


@pytest.mark.asyncio
async def test_scan_batch_limit_fails_closed_without_deactivation(
    tmp_path: Path, repository, vector_store: FakeVectorStore, embedder: FakeEmbedder
):
    from app.ingestion.service import IngestionService

    _write_policy(tmp_path / "one.docx")
    _write_policy(tmp_path / "two.docx")
    service = IngestionService(
        repository,
        vector_store,
        embedder,
        tmp_path,
        max_scan_files=1,
    )

    result = await service.scan()

    assert result.snapshot_complete is False
    assert result.deactivated == []
    assert all(item.failure_reason == "SCAN_LIMIT_EXCEEDED" for item in result.results)
    assert embedder.calls == 0


@pytest.mark.asyncio
async def test_single_file_deadline_fails_and_preserves_old_active_version(
    tmp_path: Path, repository, vector_store: FakeVectorStore
):
    from app.ingestion.service import IngestionService

    policy_path = tmp_path / "slow.docx"
    _write_policy(policy_path, "旧版虚构条款")
    first_embedder = FakeEmbedder()
    service = IngestionService(
        repository, vector_store, first_embedder, tmp_path
    )
    old = await service.index(policy_path)
    _write_policy(policy_path, "新版虚构条款")
    slow = SlowEmbedder()
    bounded = IngestionService(
        repository,
        vector_store,
        slow,
        tmp_path,
        max_file_seconds=0.001,
    )

    result = await bounded.index(policy_path, force=True)

    assert result.failure_reason == "DOCUMENT_TIMEOUT"
    assert repository.get_active(policy_path.name).version_id == old.version_id
    # The per-file deadline covers read/extract/chunk as well as provider work;
    # expiry before provider dispatch must therefore leave the provider untouched.
    assert slow.calls == 0


@pytest.mark.asyncio
async def test_racing_replacement_is_rejected_before_provider_call(
    service, repository, embedder: FakeEmbedder, policy_path: Path, tmp_path: Path, monkeypatch
):
    old = await service.index(policy_path)
    outside = tmp_path.parent / "outside.docx"
    outside.write_bytes(b"outside fictional bytes")
    original_open = os.open
    replaced = False

    def racing_open(file, flags, *args, **kwargs):
        nonlocal replaced
        if Path(file) == policy_path and not replaced:
            replaced = True
            os.replace(outside, policy_path)
        return original_open(file, flags, *args, **kwargs)

    monkeypatch.setattr("app.ingestion.service.os.open", racing_open)
    result = await service.index(policy_path, force=True)

    assert result.status == "failed"
    assert result.failure_reason == "SOURCE_CHANGED"
    assert repository.get_active(policy_path.name).version_id == old.version_id
    assert embedder.calls == 1


@pytest.mark.asyncio
async def test_hardlink_source_is_rejected_before_provider_call(
    service, embedder: FakeEmbedder, policy_path: Path, tmp_path: Path
):
    hardlink = tmp_path / "hardlink.docx"
    try:
        os.link(policy_path, hardlink)
    except OSError as error:
        pytest.skip(f"fictional hardlink capability unavailable: {error}")

    with pytest.raises(ValueError, match="safe regular file"):
        await service.index(hardlink)
    assert embedder.calls == 0


@pytest.mark.asyncio
async def test_scan_stops_at_max_plus_one_without_materializing_directory(
    tmp_path: Path, repository, vector_store: FakeVectorStore, embedder: FakeEmbedder
):
    from app.ingestion.service import IngestionService

    paths = [tmp_path / f"candidate-{index}.docx" for index in range(5)]
    for path in paths:
        _write_policy(path)
    service = IngestionService(
        repository, vector_store, embedder, tmp_path, max_scan_files=2
    )
    original_iterdir = Path.iterdir

    def bounded_iterdir(directory: Path):
        if directory.resolve() == tmp_path.resolve():
            supported_seen = 0
            for entry in original_iterdir(directory):
                if entry.suffix.casefold() in {".docx", ".pdf"}:
                    if supported_seen >= 3:
                        raise AssertionError("scan enumerated beyond max+1")
                    supported_seen += 1
                yield entry
            return
        yield from original_iterdir(directory)

    monkeypatch = pytest.MonkeyPatch()
    monkeypatch.setattr(Path, "iterdir", bounded_iterdir)
    try:
        result = await service.scan()
    finally:
        monkeypatch.undo()

    assert result.snapshot_complete is False
    assert result.deactivated == []
    assert all(item.failure_reason == "SCAN_LIMIT_EXCEEDED" for item in result.results)
    assert embedder.calls == 0


@pytest.mark.asyncio
async def test_sync_parser_timeout_fails_closed_and_preserves_old_active(
    service, repository, embedder: FakeEmbedder, policy_path: Path, monkeypatch
):
    import time

    old = await service.index(policy_path)
    service.max_file_seconds = 0.001
    from app.ingestion import service as service_module

    def blocking_parser(*args, **kwargs):
        time.sleep(0.05)
        raise AssertionError("parser worker should be cancelled before completion")

    monkeypatch.setattr(service_module, "extract_document", blocking_parser)
    result = await service.index(policy_path, force=True)

    assert result.failure_reason == "DOCUMENT_TIMEOUT"
    assert repository.get_active(policy_path.name).version_id == old.version_id
    assert embedder.calls == 1


@pytest.mark.asyncio
async def test_hardlink_added_after_open_is_rejected_before_provider_call(
    service, repository, embedder: FakeEmbedder, policy_path: Path, tmp_path: Path, monkeypatch
):
    old = await service.index(policy_path)
    original_open = os.open
    added = False

    def racing_open(file, flags, *args, **kwargs):
        nonlocal added
        descriptor = original_open(file, flags, *args, **kwargs)
        if Path(file) == policy_path and not added:
            added = True
            os.link(policy_path, tmp_path / "late-hardlink.docx")
        return descriptor

    monkeypatch.setattr("app.ingestion.service.os.open", racing_open)
    result = await service.index(policy_path, force=True)

    assert result.failure_reason == "DOCUMENT_NOT_INDEXABLE"
    assert repository.get_active(policy_path.name).version_id == old.version_id
    assert embedder.calls == 1


@pytest.mark.asyncio
async def test_sync_ocr_is_bounded_without_blocking_request_deadline(
    tmp_path: Path,
    repository,
    vector_store: FakeVectorStore,
    embedder: FakeEmbedder,
    monkeypatch: pytest.MonkeyPatch,
):
    from time import monotonic

    from app.ingestion.service import IngestionService

    policy_path = tmp_path / "slow-ocr.pdf"
    _write_blank_pdf(policy_path)
    # Keep the deadline assertion focused on the synchronous OCR adapter.
    # PyMuPDF's first render can exceed a 200 ms budget on a cold Windows
    # process, which would otherwise make this test depend on renderer startup.
    monkeypatch.setattr(
        "app.ingestion.service._render_ocr_page",
        lambda *args: b"rendered-page",
    )
    ocr = BlockingSyncOcr()
    service = IngestionService(
        repository,
        vector_store,
        embedder,
        tmp_path,
        ocr_client=ocr,
        max_file_seconds=0.2,
    )

    started = monotonic()
    result = await service.index(policy_path)
    elapsed = monotonic() - started

    assert result.failure_reason == "DOCUMENT_TIMEOUT"
    assert elapsed < 0.26
    assert ocr.calls == 1
    assert embedder.calls == 0
    assert repository.get_active(policy_path.name) is None


@pytest.mark.asyncio
async def test_pdf_render_matrix_pixel_budget_is_checked_before_pixmap(
    tmp_path: Path, repository, vector_store: FakeVectorStore, embedder: FakeEmbedder,
    monkeypatch,
):
    from app.ingestion.service import IngestionService

    policy_path = tmp_path / "giant-geometry.pdf"
    _write_blank_pdf(policy_path)
    service = IngestionService(
        repository,
        vector_store,
        embedder,
        tmp_path,
        ocr_client=FakeOcr(),
        max_pdf_pixels=20_000_000,
    )

    class Rect:
        width = 4000
        height = 4000

    class Page:
        rect = Rect()

        def get_pixmap(self, **kwargs):
            raise AssertionError("pixmap must not be allocated")

    class Pdf:
        page_count = 1

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return None

        def __iter__(self):
            return iter((Page(),))

        def load_page(self, index):
            return Page()

    monkeypatch.setattr("app.ingestion.service.fitz.open", lambda **kwargs: Pdf())

    with pytest.raises(ValueError, match="pixel limit"):
        await service._extract_with_ocr(policy_path, b"fictional pdf")


@pytest.mark.asyncio
async def test_pdf_render_dimensions_are_checked_before_ocr(
    tmp_path: Path, repository, vector_store: FakeVectorStore, embedder: FakeEmbedder,
    monkeypatch,
):
    from app.ingestion.service import IngestionService

    policy_path = tmp_path / "oversized-pixmap.pdf"
    _write_blank_pdf(policy_path)
    ocr = BlockingSyncOcr(delay=0)
    service = IngestionService(
        repository,
        vector_store,
        embedder,
        tmp_path,
        ocr_client=ocr,
        max_pdf_pixels=20_000_000,
    )

    class Rect:
        width = 100
        height = 100

    class Pixmap:
        width = 10_000
        height = 10_000

        def tobytes(self, kind):
            return b"small png"

    class Page:
        rect = Rect()

        def get_pixmap(self, **kwargs):
            return Pixmap()

    class Pdf:
        page_count = 1

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return None

        def __iter__(self):
            return iter((Page(),))

        def load_page(self, index):
            return Page()

    monkeypatch.setattr("app.ingestion.service.fitz.open", lambda **kwargs: Pdf())

    with pytest.raises(ValueError, match="pixel limit"):
        await service._extract_with_ocr(policy_path, b"fictional pdf")
    assert ocr.calls == 0


@pytest.mark.asyncio
async def test_timed_out_upsert_is_cleaned_after_worker_finishes(
    tmp_path: Path, repository, embedder: FakeEmbedder
):
    import time

    from app.ingestion.service import IngestionService

    class SlowVectorStore:
        def __init__(self):
            self.upserted: list[str] = []
            self.deleted: list[str] = []

        def upsert(self, version_id, chunks, vectors):
            time.sleep(1.0)
            self.upserted.append(version_id)

        def delete_version(self, version_id):
            self.deleted.append(version_id)

    policy_path = tmp_path / "slow-upsert.docx"
    _write_policy(policy_path)
    vector_store = SlowVectorStore()
    service = IngestionService(
        repository,
        vector_store,
        embedder,
        tmp_path,
        max_file_seconds=0.5,
    )

    result = await service.index(policy_path)
    await asyncio.sleep(1.2)

    assert result.failure_reason == "DOCUMENT_TIMEOUT"
    assert result.version_id in vector_store.deleted
    assert repository.get_active(policy_path.name) is None


@pytest.mark.asyncio
async def test_successful_upsert_with_activation_failure_keeps_old_active_version(
    service, repository, vector_store: FakeVectorStore, policy_path: Path
):
    old = await service.index(policy_path)
    _write_policy(policy_path, "虚构新版条款：交接班后复核记录。")
    connection = repository._connection
    connection.execute(
        """
        CREATE TRIGGER reject_activation
        BEFORE UPDATE OF status ON document_versions
        WHEN NEW.status = 'active'
        BEGIN
            SELECT RAISE(ABORT, 'fictional activation failure');
        END
        """
    )
    connection.commit()

    failed = await service.index(policy_path)

    assert failed.status == "failed"
    assert repository.get_active(policy_path.name).version_id == old.version_id
    assert len(vector_store.upserts) == 2


@pytest.mark.asyncio
async def test_upsert_failure_is_returned_and_partial_version_is_cleaned_without_activation(
    tmp_path: Path, repository, embedder: FakeEmbedder
):
    from app.ingestion.service import IngestionService

    policy_path = tmp_path / "partial-upsert.docx"
    _write_policy(policy_path)
    vector_store = FailingVectorStore()
    service = IngestionService(repository, vector_store, embedder, tmp_path)

    result = await service.index(policy_path)
    await service.aclose()

    assert result.status == "failed"
    assert result.failure_reason == "INDEX_FAILED"
    assert result.version_id in vector_store.deleted
    assert repository.get_active(policy_path.name) is None
    assert repository.get_version(result.version_id).status == "failed"


@pytest.mark.asyncio
async def test_cancelled_upsert_is_failed_and_drained_before_vector_close(
    tmp_path: Path, repository, embedder: FakeEmbedder
):
    from app.ingestion.service import IngestionService
    from app.main import DocumentRuntime

    events: list[str] = []
    started = threading.Event()

    class SlowVectorStore:
        def upsert(self, version_id, chunks, vectors):
            started.set()
            events.append("upsert-start")
            time.sleep(0.08)
            events.append("upsert-end")

        def delete_version(self, version_id):
            events.append("delete")

        def close(self):
            events.append("vector-close")

    class Provider:
        async def aclose(self):
            events.append("provider-close")

    class Repository:
        def close(self):
            events.append("repository-close")

    policy_path = tmp_path / "cancelled-upsert.docx"
    _write_policy(policy_path)
    vector_store = SlowVectorStore()
    service = IngestionService(
        repository,
        vector_store,
        embedder,
        tmp_path,
        max_file_seconds=1.0,
    )
    failed_versions: list[str] = []
    original_fail = repository.fail

    def record_failure(version_id: str, reason: str):
        failed_versions.append(version_id)
        original_fail(version_id, reason)

    repository.fail = record_failure
    runtime = DocumentRuntime(repository, vector_store, Provider(), service)
    task = asyncio.create_task(service.index(policy_path))
    await asyncio.to_thread(started.wait, 1.0)
    close_task = asyncio.create_task(runtime.close())
    await asyncio.sleep(0)
    assert "vector-close" not in events
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    await close_task

    assert failed_versions
    assert events.index("delete") < events.index("provider-close")
    assert events.index("delete") < events.index("vector-close")
    assert events.index("upsert-end") < events.index("vector-close")


@pytest.mark.asyncio
async def test_cleanup_failure_remains_managed_and_never_activates(
    tmp_path: Path, repository, embedder: FakeEmbedder
):
    from app.ingestion.service import IngestionService

    policy_path = tmp_path / "cleanup-failure.docx"
    _write_policy(policy_path)
    vector_store = FailingVectorStore(cleanup_fails=True)
    service = IngestionService(repository, vector_store, embedder, tmp_path)

    result = await service.index(policy_path)
    await service.aclose()

    assert result.status == "failed"
    assert repository.get_active(policy_path.name) is None
    assert vector_store.deleted == [result.version_id]
