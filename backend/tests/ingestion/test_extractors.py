import io
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import fitz
import pytest
from docx import Document

from app.ingestion.extractors import UnsupportedDocumentError, extract_document

_WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_MARKUP_COMPATIBILITY_NAMESPACE = (
    "http://schemas.openxmlformats.org/markup-compatibility/2006"
)
_WORDPROCESSING_SHAPE_NAMESPACE = (
    "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
)


def _write_fictional_textbox_docx(path: Path) -> None:
    """Create a DOCX whose only body text lives in an AlternateContent textbox."""

    source = io.BytesIO()
    Document().save(source)
    source.seek(0)
    with zipfile.ZipFile(source, "r") as archive:
        entries = {
            item.filename: archive.read(item.filename)
            for item in archive.infolist()
            if item.filename != "word/document.xml"
        }
        document_xml = archive.read("word/document.xml")

    root = ElementTree.fromstring(document_xml)
    body = root.find(f"{{{_WORD_NAMESPACE}}}body")
    assert body is not None
    alternate_content = ElementTree.fromstring(
        f"""
        <mc:AlternateContent
            xmlns:mc="{_MARKUP_COMPATIBILITY_NAMESPACE}"
            xmlns:w="{_WORD_NAMESPACE}"
            xmlns:wps="{_WORDPROCESSING_SHAPE_NAMESPACE}">
          <mc:Choice Requires="wps">
            <w:p><w:r><w:drawing><w:txbxContent>
              <w:p><w:r><w:t>文本框中的虚构制度正文。</w:t></w:r></w:p>
            </w:txbxContent></w:drawing></w:r></w:p>
          </mc:Choice>
          <mc:Fallback>
            <w:p><w:r><w:pict><w:txbxContent>
              <w:p><w:r><w:t>文本框中的虚构制度正文。</w:t></w:r></w:p>
            </w:txbxContent></w:pict></w:r></w:p>
          </mc:Fallback>
        </mc:AlternateContent>
        """
    )
    section_properties = body.find(f"{{{_WORD_NAMESPACE}}}sectPr")
    insertion_index = list(body).index(section_properties) if section_properties is not None else len(body)
    body.insert(insertion_index, alternate_content)

    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
        archive.writestr(
            "word/document.xml",
            ElementTree.tostring(root, encoding="utf-8", xml_declaration=True),
        )


def _write_explicit_heading_textbox_docx(path: Path) -> None:
    """Create a textbox document with one explicit Heading1 paragraph."""

    source = io.BytesIO()
    Document().save(source)
    source.seek(0)
    with zipfile.ZipFile(source, "r") as archive:
        entries = {
            item.filename: archive.read(item.filename)
            for item in archive.infolist()
            if item.filename != "word/document.xml"
        }
        document_xml = archive.read("word/document.xml")

    root = ElementTree.fromstring(document_xml)
    body = root.find(f"{{{_WORD_NAMESPACE}}}body")
    assert body is not None
    alternate_content = ElementTree.fromstring(
        f"""
        <mc:AlternateContent
            xmlns:mc="{_MARKUP_COMPATIBILITY_NAMESPACE}"
            xmlns:w="{_WORD_NAMESPACE}"
            xmlns:wps="{_WORDPROCESSING_SHAPE_NAMESPACE}">
          <mc:Choice Requires="wps">
            <w:p><w:r><w:drawing><w:txbxContent>
              <w:p>
                <w:pPr><w:pStyle w:val="Heading1"/></w:pPr>
                <w:r><w:t>虚构透析制度</w:t></w:r>
              </w:p>
              <w:p><w:r><w:t>护士应完成虚构设备点检。</w:t></w:r></w:p>
            </w:txbxContent></w:drawing></w:r></w:p>
          </mc:Choice>
          <mc:Fallback>
            <w:p><w:r><w:pict><w:txbxContent>
              <w:p>
                <w:pPr><w:pStyle w:val="Heading1"/></w:pPr>
                <w:r><w:t>虚构透析制度</w:t></w:r>
              </w:p>
              <w:p><w:r><w:t>护士应完成虚构设备点检。</w:t></w:r></w:p>
            </w:txbxContent></w:pict></w:r></w:p>
          </mc:Fallback>
        </mc:AlternateContent>
        """
    )
    section_properties = body.find(f"{{{_WORD_NAMESPACE}}}sectPr")
    insertion_index = list(body).index(section_properties) if section_properties is not None else len(body)
    body.insert(insertion_index, alternate_content)

    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
        archive.writestr(
            "word/document.xml",
            ElementTree.tostring(root, encoding="utf-8", xml_declaration=True),
        )


@pytest.fixture
def fictional_docx(tmp_path: Path) -> Path:
    """Create a wholly fictional policy document for real parser coverage."""

    path = tmp_path / "policy.docx"
    document = Document()
    document.add_heading("透析护理制度", level=1)
    document.add_paragraph("护士应在交接班前完成设备点检。")
    document.add_heading("培训要求", level=2)
    document.add_paragraph("新员工应参加虚构的岗前培训。")
    document.save(path)
    return path


@pytest.fixture
def fictional_pdf(tmp_path: Path) -> Path:
    path = tmp_path / "policy.pdf"
    pdf = fitz.open()
    first_page = pdf.new_page()
    first_page.insert_text(
        (72, 72), "Fictional page one: complete equipment checks before handover."
    )
    second_page = pdf.new_page()
    second_page.insert_text(
        (72, 72), "Fictional page two: record fictional induction training attendance."
    )
    pdf.save(path)
    pdf.close()
    return path


@pytest.fixture
def heading_pdf(tmp_path: Path) -> Path:
    path = tmp_path / "heading-policy.pdf"
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Fictional chapter one.")
    pdf.new_page().insert_text((72, 72), "Fictional chapter two.")
    pdf.set_toc([[1, "虚构总则", 1], [2, "设备管理", 2]])
    pdf.save(path)
    pdf.close()
    return path


@pytest.fixture
def blank_pdf(tmp_path: Path) -> Path:
    path = tmp_path / "blank.pdf"
    pdf = fitz.open()
    pdf.new_page()
    pdf.save(path)
    pdf.close()
    return path


def test_docx_keeps_heading_path_and_one_based_paragraph_source(
    fictional_docx: Path,
) -> None:
    """Breaks if heading tracking or original paragraph numbering is dropped."""

    document = extract_document(fictional_docx)

    assert document.file_name == "policy.docx"
    assert len(document.blocks) == 2
    assert document.blocks[0].source.heading_path == ("透析护理制度",)
    assert document.blocks[0].source.paragraph_start == 2
    assert document.blocks[0].source.paragraph_end == 2
    assert document.blocks[1].source.heading_path == ("透析护理制度", "培训要求")
    assert document.blocks[1].source.paragraph_start == 4


def test_docx_falls_back_to_textbox_xml_without_indexing_compatibility_duplicate(
    tmp_path: Path,
) -> None:
    path = tmp_path / "fictional-textbox.docx"
    _write_fictional_textbox_docx(path)

    document = extract_document(path)

    assert [block.text for block in document.blocks] == ["文本框中的虚构制度正文。"]
    assert document.blocks[0].source.heading_path == ("正文",)
    assert document.blocks[0].source.paragraph_start == 1


def test_docx_conservative_mode_uses_explicit_textbox_heading(tmp_path: Path) -> None:
    path = tmp_path / "explicit-heading-textbox.docx"
    _write_explicit_heading_textbox_docx(path)

    legacy = extract_document(path)
    conservative = extract_document(path, heading_mode="conservative")

    assert len(legacy.blocks) == 2
    assert all(block.source.heading_path == ("正文",) for block in legacy.blocks)
    assert [block.text for block in conservative.blocks] == ["护士应完成虚构设备点检。"]
    assert conservative.blocks[0].source.heading_path == ("虚构透析制度",)
    assert conservative.blocks[0].source.paragraph_start == 2


def test_pdf_conservative_mode_uses_explicit_bookmarks(heading_pdf: Path) -> None:
    legacy = extract_document(heading_pdf)
    conservative = extract_document(heading_pdf, heading_mode="conservative")

    assert [block.source.heading_path for block in legacy.blocks] == [
        ("正文",),
        ("正文",),
    ]
    assert [block.source.heading_path for block in conservative.blocks] == [
        ("虚构总则",),
        ("虚构总则", "设备管理"),
    ]
    assert [block.text for block in conservative.blocks] == [
        block.text for block in legacy.blocks
    ]


def test_extract_document_rejects_unknown_heading_mode(fictional_pdf: Path) -> None:
    with pytest.raises(ValueError, match="heading_mode"):
        extract_document(fictional_pdf, heading_mode="heuristic")


def test_pdf_keeps_one_based_page_source_for_extracted_text(
    fictional_pdf: Path,
) -> None:
    """Breaks if page boundaries are flattened or zero-indexed."""

    document = extract_document(fictional_pdf)

    assert document.needs_ocr is False
    assert [block.source.page for block in document.blocks] == [1, 2]
    assert (
        document.blocks[0].text
        == "Fictional page one: complete equipment checks before handover."
    )


def test_pdf_with_almost_no_text_is_marked_for_ocr(blank_pdf: Path) -> None:
    """Breaks if a scanned PDF is silently treated as an indexable text file."""

    document = extract_document(blank_pdf)

    assert document.needs_ocr is True
    assert document.blocks == ()


def test_extract_document_rejects_non_policy_file_types(tmp_path: Path) -> None:
    """Breaks if unsupported file types reach a parser."""

    path = tmp_path / "policy.txt"
    path.write_text("虚构文本", encoding="utf-8")

    with pytest.raises(UnsupportedDocumentError):
        extract_document(path)


def test_docx_limits_reject_paragraph_and_zip_expansion_before_parser(
    fictional_docx: Path,
):
    from app.ingestion.extractors import DocumentResourceLimits

    with pytest.raises(ValueError, match="paragraph"):
        extract_document(
            fictional_docx,
            limits=DocumentResourceLimits(max_docx_paragraphs=1),
        )

    expanded = fictional_docx.with_name("expanded.docx")
    expanded.write_bytes(fictional_docx.read_bytes())
    with zipfile.ZipFile(expanded, "a") as archive:
        archive.writestr("word/fictional-padding.bin", b"x" * 100)
    with pytest.raises(ValueError, match="uncompressed"):
        extract_document(
            expanded,
            limits=DocumentResourceLimits(max_docx_uncompressed_bytes=10),
        )


def test_document_resource_limits_reject_values_above_hard_bounds():
    from app.ingestion.extractors import DocumentResourceLimits

    with pytest.raises(ValueError, match="hard"):
        DocumentResourceLimits(max_file_bytes=64 * 1024 * 1024 + 1)
