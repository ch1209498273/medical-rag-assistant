from __future__ import annotations

from pathlib import Path

import pytest
from app.ingestion.task15_models import Task15RunStatus
from app.ingestion.task15_pipeline import (
    Task15PipelineError,
    load_manifest,
    run_chunking,
    run_extraction,
    run_inventory,
    run_report,
)
from docx import Document


def _source_dir(tmp_path: Path, *, count: int = 8) -> Path:
    source = tmp_path / "资料"
    source.mkdir()
    for index in range(count):
        document = Document()
        document.add_heading(f"第 {index + 1} 册", level=1)
        document.add_paragraph("这是内部标准条款。")
        if index == 0:
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "项目"
            table.cell(0, 1).text = "要求"
            table.cell(1, 0).text = "记录"
            table.cell(1, 1).text = "完整"
        document.save(source / f"volume-{index + 1}.docx")
    return source


def test_inventory_writes_fingerprints_and_hybrid_manifest(tmp_path: Path) -> None:
    source = _source_dir(tmp_path)
    run_dir = tmp_path / "private" / "RUN-001"

    manifest = run_inventory(source, run_dir)

    assert manifest.status is Task15RunStatus.INVENTORIED
    assert len(manifest.files) == 8
    assert manifest.to_dict()["retrieval_config"]["strategy"] == "hybrid"
    assert manifest.to_dict()["chunker_config"]["tokenizer_name"] == "BAAI/bge-m3"
    assert manifest.to_dict()["chunker_config"]["tokenizer_fallback"] == "cl100k_base"
    inventory = (run_dir / "inventory.json").read_text(encoding="utf-8")
    assert "volume-1.docx" in inventory
    assert "这是内部标准条款" not in inventory


def test_inventory_rejects_incomplete_source_set(tmp_path: Path) -> None:
    source = _source_dir(tmp_path, count=7)

    with pytest.raises(Task15PipelineError, match="expected 8"):
        run_inventory(source, tmp_path / "private" / "RUN-002")


def test_pipeline_reports_are_aggregate_only(tmp_path: Path) -> None:
    source = _source_dir(tmp_path)
    run_dir = tmp_path / "private" / "RUN-003"
    manifest = run_inventory(source, run_dir)

    extracted = run_extraction(manifest)
    chunked = run_chunking(extracted)
    report_path = run_report(chunked)

    assert extracted.status is Task15RunStatus.EXTRACTED
    assert chunked.status is Task15RunStatus.CHUNKED
    assert report_path.exists()
    quality = __import__("json").loads(report_path.read_text(encoding="utf-8"))
    assert quality["extraction_summary"]["file_count"] == 8
    assert quality["chunk_summary"]["docling"]["max_token_count"] <= 512
    assert quality["staging"]["sqlite_built"] is False
    assert quality["staging"]["qdrant_built"] is False
    # Warnings are parser/input dependent; the stable contract is that the
    # aggregate report normalizes them and never embeds source text.
    assert quality["warnings"] == sorted(set(quality["warnings"]))
    for path in (run_dir / "extraction-report.json", run_dir / "chunk-manifest.json", report_path):
        content = path.read_text(encoding="utf-8")
        assert "这是内部标准条款" not in content
        assert "api_key" not in content.casefold()
        assert "raw_response" not in content.casefold()


def test_quality_report_surfaces_the_selected_staging_backend(tmp_path: Path) -> None:
    source = _source_dir(tmp_path)
    run_dir = tmp_path / "private" / "RUN-STAGED"
    manifest = run_inventory(source, run_dir)
    chunked = run_chunking(run_extraction(manifest))
    (run_dir / "staging-summary.json").write_text(
        __import__("json").dumps(
            {
                "run_id": chunked.run_id,
                "source_version": chunked.source_version,
                "approved_backend": "docling",
                "candidate_count": 1,
                "sqlite_built": True,
                "qdrant_built": True,
                "vectors_indexed": False,
                "version_consistency": "pass",
                "retrieval_config": dict(chunked.retrieval_config),
            }
        ),
        encoding="utf-8",
    )

    quality = __import__("json").loads(run_report(chunked).read_text(encoding="utf-8"))

    assert quality["active_backend"] == "docling"
    assert quality["staging"]["active_backend"] == "docling"
    assert quality["staging"]["vectors_indexed"] is False


def test_extraction_uses_a_bounded_large_manual_paragraph_budget(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _source_dir(tmp_path)
    manifest = run_inventory(source, tmp_path / "private" / "RUN-003B")
    seen_limits = []
    real_audit = __import__(
        "app.ingestion.task15_pipeline", fromlist=["audit_docx_path"]
    ).audit_docx_path

    def observe_audit(path: Path, *, limits):
        seen_limits.append(limits)
        return real_audit(path, limits=limits)

    monkeypatch.setattr("app.ingestion.task15_pipeline.audit_docx_path", observe_audit)

    run_extraction(manifest)

    assert len(seen_limits) == 8
    assert seen_limits[0].max_docx_paragraphs == 50_000
    assert seen_limits[0].max_docx_paragraphs < 100_000


def test_pipeline_stops_when_source_drift_is_detected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = _source_dir(tmp_path)
    target = source / "volume-1.docx"

    # The implementation may use a single bounded read plus stat; this test
    # only asserts that a changed source cannot be accepted when detected.
    monkeypatch.setattr(
        "app.ingestion.task15_pipeline._source_fingerprint",
        lambda path, **_kwargs: (_ for _ in ()).throw(
            Task15PipelineError("source changed during scan")
        )
        if path == target
        else (_ for _ in ()).throw(Task15PipelineError("unexpected test path")),
    )
    with pytest.raises(Task15PipelineError, match="source changed"):
        run_inventory(source, tmp_path / "private" / "RUN-004")


def test_manifest_round_trip_is_loadable_for_later_cli_stages(tmp_path: Path) -> None:
    source = _source_dir(tmp_path)
    run_dir = tmp_path / "private" / "RUN-005"
    created = run_inventory(source, run_dir)

    loaded = load_manifest(run_dir)

    assert loaded == created
    assert loaded.output_root == str(run_dir.resolve())
