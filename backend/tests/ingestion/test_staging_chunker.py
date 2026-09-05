from __future__ import annotations

from pathlib import Path

import pytest
from app.ingestion.docling_adapter import StructuredDocument, StructuredElement
from app.ingestion.staging_chunker import (
    choose_active_candidate,
    chunk_with_docling,
    chunk_with_legacy_fallback,
    compare_coverage,
    validate_chunk_coverage,
)
from app.ingestion.task15_models import DocxStructureAudit, SourceFileFingerprint
from docx import Document


def _source(name: str = "manual.docx") -> SourceFileFingerprint:
    return SourceFileFingerprint(
        relative_path=name,
        size_bytes=10,
        modified_ns=20,
        sha256="a" * 64,
    )


def _document(*elements: StructuredElement) -> StructuredDocument:
    return StructuredDocument(
        source=_source(),
        elements=tuple(elements),
        parser_backend="docling",
        tokenizer_backend="fake-tokenizer",
    )


def test_docling_chunks_keep_heading_and_stable_identity() -> None:
    document = _document(
        StructuredElement(
            text="血透前核对患者身份。",
            heading_path=("第一册", "第一章", "核对"),
            content_kind="text",
            provenance_digest="b" * 64,
            paragraph_start=3,
            paragraph_end=3,
        )
    )

    first = chunk_with_docling(document, max_tokens=512)
    second = chunk_with_docling(document, max_tokens=512)

    assert first[0].record.chunk_id == second[0].record.chunk_id
    assert first[0].record.content_hash == second[0].record.content_hash
    assert first[0].record.heading_path == ("第一册", "第一章", "核对")
    assert first[0].source.heading_path == first[0].record.heading_path
    assert first[0].record.token_count > 0


def test_docling_table_chunks_repeat_header_within_token_cap() -> None:
    document = _document(
        StructuredElement(
            text="项目 | 要求\n身份 | 核对\n设备 | 检查",
            heading_path=("第一册", "表格"),
            content_kind="table",
            provenance_digest="d" * 64,
        )
    )

    candidates = chunk_with_docling(document, max_tokens=10)

    assert len(candidates) >= 2
    assert all(len(item.text) <= 10 for item in candidates)
    assert all("项目 | 要求" in item.text for item in candidates)


def test_docling_table_locator_survives_into_source_reference() -> None:
    document = _document(
        StructuredElement(
            text="项目 | 要求\n身份 | 核对",
            heading_path=("第一册", "表格"),
            content_kind="table",
            provenance_digest="c" * 64,
            table_id="table-7",
        )
    )

    candidates = chunk_with_docling(document, max_tokens=512)

    assert candidates[0].source.table_id == "table-7"


def test_docling_uses_measured_token_count_for_longer_text() -> None:
    document = _document(
        StructuredElement(
            text="条款" * 300,
            heading_path=("第一册", "第一章"),
            content_kind="text",
            provenance_digest="e" * 64,
            token_count=300,
        )
    )

    candidates = chunk_with_docling(document, max_tokens=512)

    assert len(candidates) == 1
    assert candidates[0].record.token_count == 300


def test_legacy_fallback_serializes_table_rows(tmp_path: Path) -> None:
    path = tmp_path / "manual.docx"
    document = Document()
    document.add_paragraph("表格前说明。")
    table = document.add_table(rows=4, cols=2)
    rows = [("项目", "要求"), ("身份", "核对"), ("设备", "检查"), ("记录", "完整")]
    for row, values in zip(table.rows, rows, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.save(path)

    candidates = chunk_with_legacy_fallback(path, max_tokens=10)

    table_candidates = [item for item in candidates if item.record.content_kind == "table"]
    assert table_candidates
    assert any("项目 | 要求" in item.text for item in table_candidates)
    assert all(item.record.parser_backend == "legacy" for item in candidates)


def test_coverage_report_blocks_missing_docling_table() -> None:
    audit = DocxStructureAudit(
        fingerprint=_source(),
        paragraph_count=1,
        paragraph_chars=5,
        heading_count=0,
        textbox_count=0,
        textbox_chars=0,
        table_count=1,
        nonempty_cell_count=2,
        table_chars=20,
        extracted_chars=5,
        coverage_status="blocked",
        blocking_reasons=("table_text_not_covered",),
    )
    docling = chunk_with_docling(
        _document(
            StructuredElement(
                text="正文",
                heading_path=("正文",),
                content_kind="text",
                provenance_digest="b" * 64,
            )
        )
    )

    report = compare_coverage(audit, docling_chunks=docling, fallback_chunks=())

    assert report["status"] == "blocked"
    with pytest.raises(ValueError, match="coverage"):
        validate_chunk_coverage(report)


def test_active_candidate_requires_explicit_approval() -> None:
    report = {
        "status": "pass",
        "available_backends": ["docling", "legacy"],
        "blocking_reasons": [],
    }

    with pytest.raises(ValueError, match="explicit"):
        choose_active_candidate(report)
    assert choose_active_candidate(report, approved_backend="docling") == "docling"
