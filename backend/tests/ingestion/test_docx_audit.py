from __future__ import annotations

import io
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest
from app.ingestion.docx_audit import audit_docx_bytes, audit_docx_path
from app.ingestion.extractors import DocumentResourceLimits
from docx import Document


def _docx_bytes(*, with_table: bool = False, with_textbox: bool = False) -> bytes:
    document = Document()
    document.add_heading("第一章 总则", level=1)
    document.add_paragraph("这是正文条款。")
    if with_table:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "项目"
        table.cell(0, 1).text = "要求"
        table.cell(1, 0).text = "记录"
        table.cell(1, 1).text = "完整"
    output = io.BytesIO()
    document.save(output)
    data = output.getvalue()
    if not with_textbox:
        return data

    with zipfile.ZipFile(io.BytesIO(data), "r") as source:
        members = {name: source.read(name) for name in source.namelist()}
    xml = members["word/document.xml"].decode("utf-8")
    textbox = (
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:r><w:drawing><w:txbxContent><w:p><w:r><w:t>文本框条款</w:t>'
        "</w:r></w:p></w:txbxContent></w:drawing></w:r></w:p>"
    )
    xml = xml.replace("</w:body>", textbox + "</w:body>")
    ElementTree.fromstring(xml.encode("utf-8"))
    members["word/document.xml"] = xml.encode("utf-8")
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as target:
        for name, member in members.items():
            target.writestr(name, member)
    return output.getvalue()


def test_plain_docx_audit_reports_paragraphs_and_heading(tmp_path: Path) -> None:
    path = tmp_path / "plain.docx"
    data = _docx_bytes()
    path.write_bytes(data)

    audit = audit_docx_path(path, limits=DocumentResourceLimits(max_file_bytes=64 * 1024 * 1024))

    assert audit.paragraph_count >= 1
    assert audit.heading_count == 1
    assert audit.paragraph_chars >= len("这是正文条款。")
    assert audit.coverage_status == "pass"
    assert audit.blocking_reasons == ()


def test_structured_docx_audit_blocks_uncovered_table_and_textbox(tmp_path: Path) -> None:
    path = tmp_path / "structured.docx"
    data = _docx_bytes(with_table=True, with_textbox=True)
    path.write_bytes(data)

    audit = audit_docx_bytes(
        path,
        data,
        limits=DocumentResourceLimits(max_file_bytes=64 * 1024 * 1024),
    )

    assert audit.table_count == 1
    assert audit.nonempty_cell_count == 4
    assert audit.table_chars > 0
    assert audit.textbox_count >= 1
    assert audit.textbox_chars >= len("文本框条款")
    assert audit.coverage_status == "blocked"
    assert "table_text_not_covered" in audit.blocking_reasons
    assert "textbox_text_not_covered" in audit.blocking_reasons


def test_audit_rejects_malformed_docx() -> None:
    with pytest.raises(ValueError, match="DOCX"):
        audit_docx_bytes(
            Path("broken.docx"),
            b"not a zip",
            limits=DocumentResourceLimits(max_file_bytes=64 * 1024 * 1024),
        )


def test_audit_enforces_explicit_staging_file_budget(tmp_path: Path) -> None:
    path = tmp_path / "too-large.docx"
    data = _docx_bytes()
    path.write_bytes(data)

    with pytest.raises(ValueError, match="file size"):
        audit_docx_path(path, limits=DocumentResourceLimits(max_file_bytes=len(data) - 1))
