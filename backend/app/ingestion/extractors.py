"""Text-only extractors for the supported policy formats."""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree

import fitz
from docx import Document

from app.domain.models import ExtractedBlock, ExtractedDocument, SourceRef

_HEADING_STYLE = re.compile(r"^Heading ([1-9])$", re.IGNORECASE)
_OCR_TEXT_CHARS_PER_PAGE = 20
MAX_FILE_BYTES_HARD = 64 * 1024 * 1024
MAX_PDF_PAGES_HARD = 3_000
MAX_PDF_PIXELS_HARD = 100_000_000
MAX_PDF_RENDER_BYTES_HARD = 16 * 1024 * 1024
MAX_DOCX_PARAGRAPHS_HARD = 100_000
MAX_DOCX_UNCOMPRESSED_BYTES_HARD = 256 * 1024 * 1024
MAX_DOCX_ZIP_ENTRIES_HARD = 10_000
MAX_TEXT_CHARS_HARD = 5_000_000


@dataclass(frozen=True)
class DocumentResourceLimits:
    """Finite parser budgets applied before opening document parsers."""

    max_file_bytes: int = 8 * 1024 * 1024
    max_pdf_pages: int = 200
    max_pdf_pixels: int = 20_000_000
    max_pdf_render_bytes: int = 4 * 1024 * 1024
    max_docx_paragraphs: int = 10_000
    max_docx_uncompressed_bytes: int = 32 * 1024 * 1024
    max_docx_zip_entries: int = 2_000
    max_text_chars: int = 500_000

    def __post_init__(self) -> None:
        bounds = (
            ("max_file_bytes", self.max_file_bytes, MAX_FILE_BYTES_HARD),
            ("max_pdf_pages", self.max_pdf_pages, MAX_PDF_PAGES_HARD),
            ("max_pdf_pixels", self.max_pdf_pixels, MAX_PDF_PIXELS_HARD),
            ("max_pdf_render_bytes", self.max_pdf_render_bytes, MAX_PDF_RENDER_BYTES_HARD),
            ("max_docx_paragraphs", self.max_docx_paragraphs, MAX_DOCX_PARAGRAPHS_HARD),
            ("max_docx_uncompressed_bytes", self.max_docx_uncompressed_bytes, MAX_DOCX_UNCOMPRESSED_BYTES_HARD),
            ("max_docx_zip_entries", self.max_docx_zip_entries, MAX_DOCX_ZIP_ENTRIES_HARD),
            ("max_text_chars", self.max_text_chars, MAX_TEXT_CHARS_HARD),
        )
        for name, value, maximum in bounds:
            if not isinstance(value, int) or isinstance(value, bool) or not 1 <= value <= maximum:
                raise ValueError(f"{name} exceeds the hard resource bound")


class UnsupportedDocumentError(ValueError):
    """Raised when a file is not an allowed policy document format."""


def extract_document(
    path: Path,
    *,
    data: bytes | None = None,
    limits: DocumentResourceLimits | None = None,
) -> ExtractedDocument:
    """Extract local text from a PDF or DOCX without executing document content."""

    source_path = Path(path)
    resource_limits = limits or DocumentResourceLimits()
    document_bytes = (
        _read_bounded_path(source_path, resource_limits.max_file_bytes)
        if data is None
        else data
    )
    if len(document_bytes) > resource_limits.max_file_bytes:
        raise ValueError("document file size limit exceeded")
    suffix = source_path.suffix.casefold()
    if suffix == ".docx":
        return _extract_docx(source_path, document_bytes, resource_limits)
    if suffix == ".pdf":
        return _extract_pdf(source_path, document_bytes, resource_limits)
    raise UnsupportedDocumentError(f"Unsupported document type: {source_path.suffix}")


def _extract_docx(
    path: Path, data: bytes, limits: DocumentResourceLimits
) -> ExtractedDocument:
    _check_docx_zip(data, limits)
    document = Document(io.BytesIO(data))
    if len(document.paragraphs) > limits.max_docx_paragraphs:
        raise ValueError("DOCX paragraph limit exceeded")
    headings: list[str] = []
    blocks: list[ExtractedBlock] = []
    total_characters = 0
    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        total_characters += len(text)
        if total_characters > limits.max_text_chars:
            raise ValueError("document text character limit exceeded")
        heading_match = _HEADING_STYLE.match(paragraph.style.name)
        if heading_match:
            if text:
                level = int(heading_match.group(1))
                headings = headings[: level - 1]
                headings.append(text)
            continue
        if text:
            blocks.append(
                ExtractedBlock(
                    text=text,
                    source=SourceRef(
                        file_name=path.name,
                        heading_path=tuple(headings) or ("正文",),
                        paragraph_start=paragraph_number,
                        paragraph_end=paragraph_number,
                    ),
                )
            )
    if not blocks:
        blocks = _extract_docx_xml_textboxes(path, data, limits)
    return ExtractedDocument(file_name=path.name, blocks=tuple(blocks))


def _extract_docx_xml_textboxes(
    path: Path, data: bytes, limits: DocumentResourceLimits
) -> list[ExtractedBlock]:
    """Recover text from drawing textboxes ignored by ``python-docx``.

    WPS and Word can store the same textbox twice inside ``mc:Choice`` and
    ``mc:Fallback`` compatibility branches.  We read the preferred branch and
    skip the fallback so one policy paragraph is not indexed twice.
    """

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            document_xml = archive.read("word/document.xml")
        root = ElementTree.fromstring(document_xml)
    except (KeyError, zipfile.BadZipFile, ElementTree.ParseError) as error:
        raise ValueError("DOCX XML is invalid") from error

    paragraphs: list[str] = []
    total_characters = 0

    def walk(element: ElementTree.Element) -> None:
        nonlocal total_characters
        local_name = element.tag.rsplit("}", 1)[-1]
        if local_name in {"Fallback", "del"}:
            return
        if local_name == "p":
            text = _docx_paragraph_own_text(element)
            if text:
                paragraphs.append(text)
                total_characters += len(text)
                if total_characters > limits.max_text_chars:
                    raise ValueError("document text character limit exceeded")
        for child in element:
            walk(child)

    walk(root)
    return [
        ExtractedBlock(
            text=text,
            source=SourceRef(
                file_name=path.name,
                heading_path=("正文",),
                paragraph_start=paragraph_number,
                paragraph_end=paragraph_number,
            ),
        )
        for paragraph_number, text in enumerate(paragraphs, start=1)
    ]


def _docx_paragraph_own_text(paragraph: ElementTree.Element) -> str:
    """Read one paragraph without absorbing text from nested textbox paragraphs."""

    text_parts: list[str] = []

    def collect(element: ElementTree.Element, *, root: bool = False) -> None:
        local_name = element.tag.rsplit("}", 1)[-1]
        if local_name in {"Fallback", "del"}:
            return
        if not root and local_name == "p":
            return
        if local_name == "t" and element.text:
            text_parts.append(element.text)
        for child in element:
            collect(child)

    collect(paragraph, root=True)
    return "".join(text_parts).strip()


def _extract_pdf(
    path: Path, data: bytes, limits: DocumentResourceLimits
) -> ExtractedDocument:
    blocks: list[ExtractedBlock] = []
    total_characters = 0
    with fitz.open(stream=data, filetype="pdf") as document:
        page_count = document.page_count
        if page_count > limits.max_pdf_pages:
            raise ValueError("PDF page limit exceeded")
        for page_number, page in enumerate(document, start=1):
            if page.rect.width * page.rect.height > limits.max_pdf_pixels:
                raise ValueError("PDF page pixel limit exceeded")
            text = page.get_text("text").strip()
            total_characters += len(text)
            if total_characters > limits.max_text_chars:
                raise ValueError("document text character limit exceeded")
            if text:
                blocks.append(
                    ExtractedBlock(
                        text=text,
                        source=SourceRef(
                            file_name=path.name,
                            heading_path=("正文",),
                            page=page_number,
                            page_end=page_number,
                        ),
                    )
                )
    needs_ocr = page_count == 0 or total_characters / page_count < _OCR_TEXT_CHARS_PER_PAGE
    return ExtractedDocument(
        file_name=path.name,
        blocks=tuple(blocks),
        needs_ocr=needs_ocr,
    )


def _check_docx_zip(data: bytes, limits: DocumentResourceLimits) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > limits.max_docx_zip_entries:
                raise ValueError("DOCX ZIP entry limit exceeded")
            total_uncompressed = 0
            for entry in entries:
                if entry.flag_bits & 0x1:
                    raise ValueError("encrypted DOCX ZIP entries are not accepted")
                total_uncompressed += entry.file_size
                if total_uncompressed > limits.max_docx_uncompressed_bytes:
                    raise ValueError("DOCX uncompressed size limit exceeded")
            _check_docx_xml_limits(archive, limits)
    except zipfile.BadZipFile as error:
        raise ValueError("DOCX archive is invalid") from error


def _read_bounded_path(path: Path, maximum: int) -> bytes:
    with path.open("rb") as handle:
        data = handle.read(maximum + 1)
    if len(data) > maximum:
        raise ValueError("document file size limit exceeded")
    return data


def _check_docx_xml_limits(
    archive: zipfile.ZipFile, limits: DocumentResourceLimits
) -> None:
    paragraph_count = 0
    text_characters = 0
    for entry in archive.infolist():
        if not entry.filename.startswith("word/") or not entry.filename.endswith(".xml"):
            continue
        with archive.open(entry, "r") as stream:
            try:
                for _, element in ElementTree.iterparse(stream, events=("end",)):
                    tag = element.tag.rsplit("}", 1)[-1]
                    if tag == "p":
                        paragraph_count += 1
                        if paragraph_count > limits.max_docx_paragraphs:
                            raise ValueError("DOCX paragraph limit exceeded")
                    if tag == "t" and element.text:
                        text_characters += len(element.text)
                        if text_characters > limits.max_text_chars:
                            raise ValueError("document text character limit exceeded")
                    element.clear()
            except ElementTree.ParseError as error:
                raise ValueError("DOCX XML is invalid") from error
