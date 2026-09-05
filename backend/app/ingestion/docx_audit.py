"""Read-only DOCX structure and extraction-coverage auditing."""

from __future__ import annotations

import io
import re
import zipfile
from hashlib import sha256
from pathlib import Path
from xml.etree import ElementTree

from docx import Document

from app.ingestion.extractors import (
    DocumentResourceLimits,
    _check_docx_zip,
    extract_document,
)
from app.ingestion.task15_models import DocxStructureAudit, SourceFileFingerprint

_HEADING_STYLE = re.compile(r"^Heading ([1-9])$", re.IGNORECASE)
_DEFAULT_LIMITS = DocumentResourceLimits(
    max_file_bytes=64 * 1024 * 1024,
    max_docx_uncompressed_bytes=64 * 1024 * 1024,
)


def audit_docx_path(
    path: Path,
    *,
    limits: DocumentResourceLimits | None = None,
) -> DocxStructureAudit:
    """Read one DOCX with a bounded read and return aggregate structure facts."""

    source_path = Path(path)
    resource_limits = limits or _DEFAULT_LIMITS
    try:
        with source_path.open("rb") as handle:
            data = handle.read(resource_limits.max_file_bytes + 1)
    except OSError as error:
        raise ValueError("DOCX source cannot be read") from error
    if len(data) > resource_limits.max_file_bytes:
        raise ValueError("document file size limit exceeded")
    return audit_docx_bytes(source_path, data, limits=resource_limits)


def audit_docx_bytes(
    path: Path,
    data: bytes,
    *,
    limits: DocumentResourceLimits | None = None,
) -> DocxStructureAudit:
    """Audit paragraphs, headings, textboxes, tables and extracted coverage.

    The returned object contains counts and digests only.  It intentionally
    does not return document text, so it can safely be included in a private
    aggregate report.
    """

    source_path = Path(path)
    resource_limits = limits or _DEFAULT_LIMITS
    if not isinstance(data, bytes):
        raise TypeError("DOCX data must be bytes")
    if len(data) > resource_limits.max_file_bytes:
        raise ValueError("document file size limit exceeded")
    _check_docx_zip(data, resource_limits)

    try:
        document = Document(io.BytesIO(data))
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            document_xml = archive.read("word/document.xml")
        root = ElementTree.fromstring(document_xml)
    except (KeyError, zipfile.BadZipFile, ElementTree.ParseError, ValueError) as error:
        raise ValueError("DOCX document is invalid") from error
    except Exception as error:
        raise ValueError("DOCX document is invalid") from error

    non_heading_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if _HEADING_STYLE.match(paragraph.style.name or "") is None
    ]
    paragraph_chars = sum(len(paragraph.text.strip()) for paragraph in non_heading_paragraphs)
    heading_count = sum(
        1
        for paragraph in document.paragraphs
        if _HEADING_STYLE.match(paragraph.style.name or "") is not None
        and paragraph.text.strip()
    )

    tables = list(document.tables)
    nonempty_cell_count = 0
    table_chars = 0
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    nonempty_cell_count += 1
                    table_chars += len(text)

    textbox_count, textbox_chars = _textbox_facts(root)
    extracted = extract_document(
        source_path,
        data=data,
        limits=resource_limits,
        heading_mode="conservative",
    )
    extracted_chars = sum(len(block.text) for block in extracted.blocks)

    reasons: list[str] = []
    if extracted_chars < paragraph_chars:
        reasons.append("paragraph_text_not_covered")
    expected_with_tables = paragraph_chars + table_chars
    if table_chars and extracted_chars < expected_with_tables:
        reasons.append("table_text_not_covered")
    expected_with_textboxes = expected_with_tables + textbox_chars
    if textbox_chars and extracted_chars < expected_with_textboxes:
        reasons.append("textbox_text_not_covered")
    coverage_status = "blocked" if reasons else "pass"

    try:
        modified_ns = source_path.stat().st_mtime_ns
    except OSError:
        modified_ns = 0
    fingerprint = SourceFileFingerprint(
        relative_path=source_path.name,
        size_bytes=len(data),
        modified_ns=modified_ns,
        sha256=sha256(data).hexdigest(),
    )
    return DocxStructureAudit(
        fingerprint=fingerprint,
        paragraph_count=len(document.paragraphs),
        paragraph_chars=paragraph_chars,
        heading_count=heading_count,
        textbox_count=textbox_count,
        textbox_chars=textbox_chars,
        table_count=len(tables),
        nonempty_cell_count=nonempty_cell_count,
        table_chars=table_chars,
        extracted_chars=extracted_chars,
        coverage_status=coverage_status,
        blocking_reasons=tuple(reasons),
    )


def explain_coverage(audit: DocxStructureAudit) -> tuple[str, ...]:
    """Return stable, human-readable coverage reasons without exposing text."""

    if audit.coverage_status == "pass":
        return ()
    return audit.blocking_reasons


def _textbox_facts(root: ElementTree.Element) -> tuple[int, int]:
    count = 0
    characters = 0
    for element in root.iter():
        if element.tag.rsplit("}", 1)[-1] != "txbxContent":
            continue
        text = "".join(
            child.text or ""
            for child in element.iter()
            if child.tag.rsplit("}", 1)[-1] == "t"
        ).strip()
        if text:
            count += 1
            characters += len(text)
    return count, characters


__all__ = ["audit_docx_bytes", "audit_docx_path", "explain_coverage"]
